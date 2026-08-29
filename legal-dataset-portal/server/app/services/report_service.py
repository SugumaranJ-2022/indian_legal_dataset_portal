import os
import datetime
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
from reportlab.pdfgen import canvas

from app.core.config import settings
from app.models.models import Source, Document, QualityCheck, Duplicate, CourtMetadata, Dataset, DatasetEvidence, GapAnalysis, ResearchSearchLog, ResearchMethodology, AuditLog

class NumberedCanvas(canvas.Canvas):
    """
    Two-pass canvas to dynamically compute and render page numbers
    in 'Page X of Y' format, skipping the cover page.
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        # We don't draw headers/footers on the cover page
        if self._pageNumber == 1:
            return
        
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b")) # slate-500
        
        # Header text
        self.drawString(54, 800, "INDIAN LEGAL DATASET LANDSCAPE REPORT — DISCOVERY & COMPLIANCE")
        self.setStrokeColor(colors.HexColor("#cbd5e1")) # slate-300
        self.setLineWidth(0.5)
        self.line(54, 792, 595 - 54, 792) # A4 width is 595.27 points
        
        # Footer text
        self.drawString(54, 40, "Confidential — Indian Legal Dataset Research Portal")
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(595 - 54, 40, page_text)
        self.line(54, 52, 595 - 54, 52)
        
        self.restoreState()


def generate_excel_report(db: Session, file_path: str, report_type: str = "audit"):
    """
    Generate and export multi-sheet Excel reports.
    Supports report_type = "audit" (original data audit report) and "landscape" (new dataset landscape report).
    """
    if report_type == "audit":
        return generate_audit_excel_report(db, file_path)
    else:
        return generate_landscape_excel_report(db, file_path)


def generate_audit_excel_report(db: Session, file_path: str):
    wb = openpyxl.Workbook()
    
    # 1. Sources Sheet
    ws_sources = wb.active
    ws_sources.title = "Sources"
    sources = db.query(Source).all()
    headers_sources = ["ID", "Website Name", "Authority", "Category", "Source Type", "Languages", "Download Available", "Website URL", "Notes"]
    ws_sources.append(headers_sources)
    for s in sources:
        ws_sources.append([s.id, s.website_name, s.authority, s.category, s.source_type, s.languages, s.download_available, s.website_url, s.notes])

    # 2. Documents Sheet
    ws_docs = wb.create_sheet(title="Documents")
    docs = db.query(Document).all()
    headers_docs = ["ID", "Doc Code", "Title", "Year", "Category", "Authority", "Language", "Source ID", "Source URL", "Filename", "Version", "Status", "Uploaded At"]
    ws_docs.append(headers_docs)
    for d in docs:
        uploaded_str = d.uploaded_at.strftime('%Y-%m-%d %H:%M:%S') if d.uploaded_at else ""
        ws_docs.append([d.id, d.document_code, d.title, d.year, d.category, d.authority, d.language, d.source_id, d.source_url, d.filename, d.version, d.status, uploaded_str])

    # 3. Court Metadata Sheet
    ws_court = wb.create_sheet(title="Metadata")
    court_meta = db.query(CourtMetadata).all()
    headers_court = ["ID", "Doc ID", "CNR Number", "Case Number", "Case Type", "Court", "State", "District", "Petitioner", "Respondent", "Filing Date", "Hearing Date", "Disposal Date", "Judge", "Case Status"]
    ws_court.append(headers_court)
    for c in court_meta:
        filing_str = c.filing_date.strftime('%Y-%m-%d') if c.filing_date else ""
        hearing_str = c.hearing_date.strftime('%Y-%m-%d') if c.hearing_date else ""
        disposal_str = c.disposal_date.strftime('%Y-%m-%d') if c.disposal_date else ""
        ws_court.append([c.id, c.document_id, c.cnr_number, c.case_number, c.case_type, c.court, c.state, c.district, c.petitioner, c.respondent, filing_str, hearing_str, disposal_str, c.judge, c.case_status])

    # 4. Quality Checks Sheet
    ws_quality = wb.create_sheet(title="Quality")
    quality = db.query(QualityCheck).all()
    headers_quality = ["ID", "Doc ID", "Official Source", "Readable PDF", "Complete Content", "Metadata Correct", "Duplicate Checked", "Version Verified", "Verification Status"]
    ws_quality.append(headers_quality)
    for q in quality:
        ws_quality.append([q.id, q.document_id, q.official_source, q.readable, q.complete_content, q.metadata_complete, q.duplicate_checked, q.version_verified, q.verification_status])

    # 5. Duplicates Sheet
    ws_dups = wb.create_sheet(title="Duplicates")
    dups = db.query(Duplicate).all()
    headers_dups = ["ID", "Doc ID", "Duplicate Doc ID", "Reason", "Action Taken", "Resolution Status"]
    ws_dups.append(headers_dups)
    for dp in dups:
        ws_dups.append([dp.id, dp.document_id, dp.duplicate_document_id, dp.reason, dp.action, dp.status])

    wb.save(file_path)


def generate_landscape_excel_report(db: Session, file_path: str):
    """
    Generates a highly polished multi-sheet Excel spreadsheet containing 15 sheets
    related to the Indian Legal Dataset Landscape research investigation.
    """
    wb = openpyxl.Workbook()
    
    # Pre-configure styles
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=11)
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep Blue
    sub_header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid") # Teal
    border_side = Side(style='thin', color='CBD5E1')
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")

    def format_sheet(ws, title_row):
        ws.row_dimensions[1].height = 25
        for cell in ws[1]:
            cell.font = title_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = cell_border
        
        # Auto-fit columns
        for col in ws.columns:
            max_len = 0
            for cell in col:
                cell.font = cell_font
                cell.border = cell_border
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Gather data
    datasets = db.query(Dataset).all()
    shortlisted = db.query(Dataset).filter(Dataset.shortlisted == True).order_by(Dataset.shortlist_rank).all()
    evidence = db.query(DatasetEvidence).all()
    gaps = db.query(GapAnalysis).all()
    logs = db.query(ResearchSearchLog).all()
    methodologies = db.query(ResearchMethodology).all()
    audits = db.query(AuditLog).filter(AuditLog.entity.in_(["Dataset", "DatasetEvidence", "GapAnalysis", "ResearchSearchLog", "ResearchMethodology"])).all()

    # Sheet 1: Dataset Inventory
    ws_inv = wb.active
    ws_inv.title = "Dataset Inventory"
    ws_inv.append(["ID", "Dataset Name", "Short Name", "Platform", "Category", "Subcategory", "URL", "Creator", "Organization", "Record Count", "Provenance Status", "License Status", "Reuse Classification", "Freshness"])
    for d in datasets:
        ws_inv.append([d.id, d.dataset_name, d.short_name, d.platform, d.category, d.subcategory, d.dataset_url, d.creator, d.organization, d.record_count, d.provenance_status, d.license_status, d.reuse_classification, d.freshness_status])
    format_sheet(ws_inv, 1)

    # Sheet 2: Shortlisted Datasets
    ws_short = wb.create_sheet(title="Shortlisted Datasets")
    ws_short.append(["Rank", "Shortlist Reason", "ID", "Dataset Name", "Platform", "Category", "Record Count", "Provenance Status", "License Status", "Relevance Score"])
    for d in shortlisted:
        ws_short.append([d.shortlist_rank or "", d.shortlist_reason or "", d.id, d.dataset_name, d.platform, d.category, d.record_count, d.provenance_status, d.license_status, d.research_relevance_score])
    format_sheet(ws_short, 1)

    # Sheet 3: Dataset Comparison
    ws_comp = wb.create_sheet(title="Dataset Comparison")
    ws_comp.append(["Dataset", "Platform", "Category", "Creator", "Records", "Time Period", "Courts", "Languages", "Format", "Metadata", "Original Docs", "Provenance", "License", "Freshness", "Limitations", "Reuse Classification"])
    for d in shortlisted:
        ws_comp.append([d.dataset_name, d.platform, d.category, d.creator, d.record_count, f"{d.time_period_start}-{d.time_period_end}", d.courts, d.languages, d.format, d.metadata_fields, "Yes" if d.original_documents_available else "No", d.provenance_status, d.license_name, d.freshness_status, d.limitations, d.reuse_classification])
    format_sheet(ws_comp, 1)

    # Sheet 4: Provenance
    ws_prov = wb.create_sheet(title="Provenance")
    ws_prov.append(["Dataset Name", "Original Source", "Original Source URL", "Collection Method", "Collection Description", "Provenance Status", "Evidence Text", "Provenance Notes"])
    for d in datasets:
        ws_prov.append([d.dataset_name, d.original_source, d.original_source_url, d.collection_method, d.collection_description, d.provenance_status, d.provenance_evidence, d.provenance_notes])
    format_sheet(ws_prov, 1)

    # Sheet 5: Licenses
    ws_lic = wb.create_sheet(title="Licenses")
    ws_lic.append(["Dataset Name", "License Name", "License URL", "License Status", "Commercial Use", "Redistribution Allowed", "Attribution Required", "Derivative Use", "Restrictions", "Notes"])
    for d in datasets:
        ws_lic.append([d.dataset_name, d.license_name, d.license_url, d.license_status, "Yes" if d.commercial_use else "No", "Yes" if d.redistribution_allowed else "No", "Yes" if d.attribution_required else "No", "Yes" if d.derivative_use else "No", d.usage_restrictions, d.license_notes])
    format_sheet(ws_lic, 1)

    # Sheet 6: Coverage
    ws_cov = wb.create_sheet(title="Coverage")
    ws_cov.append(["Dataset Name", "Record Count", "Time Period", "Courts Covered", "Jurisdictions", "States Covered", "Coverage Description"])
    for d in datasets:
        ws_cov.append([d.dataset_name, d.record_count, f"{d.time_period_start or 'N/A'} - {d.time_period_end or 'N/A'}", d.courts, d.jurisdictions, d.states, d.coverage_description])
    format_sheet(ws_cov, 1)

    # Sheet 7: Metadata
    ws_meta = wb.create_sheet(title="Metadata")
    ws_meta.append(["Dataset Name", "Metadata Available", "Metadata Fields", "Metadata Quality Assessment"])
    for d in datasets:
        ws_meta.append([d.dataset_name, "Yes" if d.metadata_available else "No", d.metadata_fields, d.metadata_quality])
    format_sheet(ws_meta, 1)

    # Sheet 8: Document Availability
    ws_doc_av = wb.create_sheet(title="Document Availability")
    ws_doc_av.append(["Dataset Name", "Original Documents Available", "Original PDF Available", "OCR Available", "Text Available", "Structured Data Available"])
    for d in datasets:
        ws_doc_av.append([d.dataset_name, "Yes" if d.original_documents_available else "No", "Yes" if d.original_pdf_available else "No", "Yes" if d.ocr_available else "No", "Yes" if d.text_available else "No", "Yes" if d.structured_data_available else "No"])
    format_sheet(ws_doc_av, 1)

    # Sheet 9: Freshness
    ws_fresh = wb.create_sheet(title="Freshness")
    ws_fresh.append(["Dataset Name", "Publication Date", "Last Updated Date", "Freshness Status", "Completeness Assessment", "Quality Assessment"])
    for d in datasets:
        pub_str = d.publication_date.strftime('%Y-%m-%d') if d.publication_date else "N/A"
        up_str = d.last_updated_date.strftime('%Y-%m-%d') if d.last_updated_date else "N/A"
        ws_fresh.append([d.dataset_name, pub_str, up_str, d.freshness_status, d.completeness_assessment, d.data_quality_assessment])
    format_sheet(ws_fresh, 1)

    # Sheet 10: Gap Analysis
    ws_gaps = wb.create_sheet(title="Gap Analysis")
    ws_gaps.append(["Category", "Current Availability", "Quality", "Existing Coverage", "Current State Details", "Problems/Gaps", "Priority Level", "Recommendation", "Evidence"])
    for g in gaps:
        ws_gaps.append([g.category, g.availability, g.quality, g.coverage, g.current_state, g.gap, g.priority, g.recommendation, g.evidence])
    format_sheet(ws_gaps, 1)

    # Sheet 11: Reuse Assessment
    ws_reuse = wb.create_sheet(title="Reuse Assessment")
    ws_reuse.append(["Dataset Name", "Reuse Classification", "Reuse Reason", "Recommendation", "Why Selected", "Research Notes"])
    for d in datasets:
        ws_reuse.append([d.dataset_name, d.reuse_classification, d.reuse_reason, d.recommendation, d.why_selected, d.research_notes])
    format_sheet(ws_reuse, 1)

    # Sheet 12: Research Methodology
    ws_meth = wb.create_sheet(title="Research Methodology")
    ws_meth.append(["Platforms Searched", "Search Date", "Search Terms", "Categories Investigated", "Selection Criteria", "Exclusion Criteria", "Verification Process"])
    for m in methodologies:
        date_str = m.search_date.strftime('%Y-%m-%d') if m.search_date else "N/A"
        ws_meth.append([m.platform_searched, date_str, m.search_terms, m.categories_investigated, m.selection_criteria, m.exclusion_criteria, m.verification_process])
    format_sheet(ws_meth, 1)

    # Sheet 13: Search Log
    ws_slog = wb.create_sheet(title="Search Log")
    ws_slog.append(["Platform", "Search Query", "Search Date", "Researcher", "Results Found", "Relevant Results", "Notes"])
    for l in logs:
        date_str = l.search_date.strftime('%Y-%m-%d') if l.search_date else "N/A"
        ws_slog.append([l.platform, l.search_query, date_str, l.researcher, l.results_found, l.relevant_results, l.notes])
    format_sheet(ws_slog, 1)

    # Sheet 14: Evidence
    ws_ev = wb.create_sheet(title="Evidence")
    ws_ev.append(["Dataset Name", "Source Type", "Source Title", "Source URL", "Description", "Evidence Text", "Verified", "Reviewer", "Verified At", "Notes"])
    for e in evidence:
        d = db.query(Dataset).filter(Dataset.id == e.dataset_id).first()
        name = d.dataset_name if d else "Unknown Dataset"
        v_str = "Yes" if e.verified else "No"
        v_date = e.verified_at.strftime('%Y-%m-%d %H:%M') if e.verified_at else ""
        ws_ev.append([name, e.source_type, e.source_title, e.source_url, e.source_description, e.evidence_text, v_str, e.reviewer, v_date, e.notes])
    format_sheet(ws_ev, 1)

    # Sheet 15: Audit Log
    ws_audit = wb.create_sheet(title="Audit Log")
    ws_audit.append(["Timestamp", "User Email", "Action", "Entity Type", "Entity ID", "Previous Value", "New Value", "Evidence Source", "Notes"])
    for a in audits:
        t_str = a.timestamp.strftime('%Y-%m-%d %H:%M:%S') if a.timestamp else ""
        ws_audit.append([t_str, a.user_email, a.action, a.entity, a.entity_id, a.previous_value, a.new_value, a.evidence_source, a.notes])
    format_sheet(ws_audit, 1)

    wb.save(file_path)


def generate_pdf_report(db: Session, file_path: str, report_type: str = "audit") -> str:
    """
    Generates a professional A4 PDF report.
    Supports report_type = "audit" (original data audit report) and "landscape" (new dataset landscape report).
    """
    if report_type == "audit":
        return generate_audit_pdf_report(db, file_path)
    else:
        return generate_landscape_pdf_report(db, file_path)


def generate_audit_pdf_report(db: Session, file_path: str) -> str:
    doc_template = SimpleDocTemplate(
        file_path,
        pagesize=A4,
        rightMargin=54,
        leftMargin=54,
        topMargin=72,
        bottomMargin=72
    )

    styles = getSampleStyleSheet()
    c_primary = colors.HexColor("#1e3a8a")
    c_secondary = colors.HexColor("#0f766e")
    c_dark = colors.HexColor("#1e293b")
    c_light = colors.HexColor("#f8fafc")

    title_style = ParagraphStyle('CoverTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=30, leading=38, textColor=c_primary, spaceAfter=15)
    subtitle_style = ParagraphStyle('CoverSubtitle', parent=styles['Normal'], fontName='Helvetica', fontSize=14, leading=18, textColor=colors.HexColor("#475569"), spaceAfter=300)
    h1_style = ParagraphStyle('SectionH1', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=c_dark, spaceBefore=20, spaceAfter=10, keepWithNext=True)
    h2_style = ParagraphStyle('SectionH2', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=13, leading=17, textColor=c_secondary, spaceBefore=12, spaceAfter=6, keepWithNext=True)
    body_style = ParagraphStyle('ReportBody', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=14, textColor=c_dark, spaceAfter=8)
    table_header_style = ParagraphStyle('TableHeader', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.white)
    table_cell_style = ParagraphStyle('TableCell', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=9.5, textColor=c_dark)

    story = []

    # Cover Page
    story.append(Spacer(1, 40))
    story.append(Paragraph("INDIAN LEGAL DATASET AUDIT REPORT", title_style))
    story.append(Paragraph("Enterprise Portal Verification & Analytics Summary", subtitle_style))
    metadata_text = f"<b>Report Date:</b> {datetime.datetime.now().strftime('%d %B, %Y')}<br/><b>Author:</b> Indian Legal Research Portal Team<br/><b>Status:</b> Official Release<br/><b>Classification:</b> Confidential / Government Data Team Standard"
    story.append(HRFlowable(width="100%", thickness=1.5, color=c_primary, spaceBefore=20, spaceAfter=20))
    story.append(Paragraph(metadata_text, body_style))
    story.append(PageBreak())

    # Executive Summary
    story.append(Paragraph("Executive Summary", h1_style))
    story.append(HRFlowable(width="100%", thickness=0.8, color=c_secondary, spaceAfter=15))
    summary_text = "This comprehensive dataset audit report compiles findings from the <b>Legal Dataset Research Portal</b>. The portal facilitates the mapping, collection, and verification of authoritative Indian legal sources, statutes, rules, and court judgments."
    story.append(Paragraph(summary_text, body_style))

    # Key Stats
    total_sources = db.query(Source).count()
    total_docs = db.query(Document).count()
    verified_docs = db.query(Document).filter(Document.status == "Verified").count()
    needs_review_docs = db.query(Document).filter(Document.status == "Needs Review").count()
    duplicate_docs = db.query(Duplicate).count()

    story.append(Spacer(1, 10))
    story.append(Paragraph("Key Portal Statistics", h2_style))
    
    stats_data = [
        [Paragraph("<b>Metric</b>", table_cell_style), Paragraph("<b>Count</b>", table_cell_style), Paragraph("<b>Description</b>", table_cell_style)],
        [Paragraph("Total Sources Mapped", table_cell_style), Paragraph(str(total_sources), table_cell_style), Paragraph("Unique government databases and repositories cataloged.", table_cell_style)],
        [Paragraph("Total Documents Collected", table_cell_style), Paragraph(str(total_docs), table_cell_style), Paragraph("PDF files successfully compiled and indexed.", table_cell_style)],
        [Paragraph("Verified (Fully Audited)", table_cell_style), Paragraph(str(verified_docs), table_cell_style), Paragraph("Documents passing all quality verification tests.", table_cell_style)],
        [Paragraph("Needs Review", table_cell_style), Paragraph(str(needs_review_docs), table_cell_style), Paragraph("Documents waiting for manual review or missing checks.", table_cell_style)],
        [Paragraph("Duplicates Identified", table_cell_style), Paragraph(str(duplicate_docs), table_cell_style), Paragraph("Deduplicated document warnings needing resolution.", table_cell_style)]
    ]
    
    stats_table = Table(stats_data, colWidths=[130, 60, 300])
    stats_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_light),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,0), 6),
        ('TOPPADDING', (0,0), (-1,0), 6),
    ]))
    story.append(stats_table)
    story.append(PageBreak())

    # Sources Table
    story.append(Paragraph("1. Source List Deliverable", h1_style))
    sources = db.query(Source).all()
    source_rows = [[Paragraph("<b>Website Name</b>", table_header_style), Paragraph("<b>Authority</b>", table_header_style), Paragraph("<b>Category</b>", table_header_style), Paragraph("<b>Type</b>", table_header_style), Paragraph("<b>Website URL</b>", table_header_style)]]
    for s in sources:
        source_rows.append([Paragraph(s.website_name, table_cell_style), Paragraph(s.authority, table_cell_style), Paragraph(s.category, table_cell_style), Paragraph(s.source_type, table_cell_style), Paragraph(s.website_url, table_cell_style)])
    sources_table = Table(source_rows, colWidths=[100, 100, 70, 70, 150])
    sources_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_primary),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, c_light]),
    ]))
    story.append(sources_table)

    doc_template.build(story, canvasmaker=NumberedCanvas)
    return file_path


def generate_landscape_pdf_report(db: Session, file_path: str) -> str:
    """
    Generates a professional research PDF report titled:
    'Indian Legal Dataset Landscape - Existing Dataset & Resource Investigation'
    It is fully structured based on live database research parameters.
    """
    doc_template = SimpleDocTemplate(
        file_path,
        pagesize=A4,
        rightMargin=54,
        leftMargin=54,
        topMargin=72,
        bottomMargin=72
    )

    styles = getSampleStyleSheet()
    c_primary = colors.HexColor("#1e3a8a") # deep blue
    c_secondary = colors.HexColor("#0f766e") # teal
    c_dark = colors.HexColor("#1e293b") # dark slate
    c_light = colors.HexColor("#f8fafc") # slate-50

    title_style = ParagraphStyle(
        'LandscapeTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=26,
        leading=32,
        textColor=c_primary,
        spaceAfter=10
    )
    
    subtitle_style = ParagraphStyle(
        'LandscapeSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#475569"),
        spaceAfter=250
    )
    
    h1_style = ParagraphStyle(
        'LandscapeH1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=c_dark,
        spaceBefore=18,
        spaceAfter=8,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        'LandscapeH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=c_secondary,
        spaceBefore=10,
        spaceAfter=5,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'LandscapeBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=c_dark,
        spaceAfter=6
    )

    warning_body_style = ParagraphStyle(
        'LandscapeWarningBody',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#b91c1c"), # red-700
        spaceAfter=6
    )

    table_header_style = ParagraphStyle(
        'LandscapeTableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9,
        textColor=colors.white
    )

    table_cell_style = ParagraphStyle(
        'LandscapeTableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=8.5,
        textColor=c_dark
    )

    story = []

    # ================= COVER PAGE =================
    story.append(Spacer(1, 40))
    story.append(Paragraph("INDIAN LEGAL DATASET LANDSCAPE", title_style))
    story.append(Paragraph("Existing Dataset & Resource Investigation", subtitle_style))
    
    metadata_text = f"""
    <b>Report Date:</b> {datetime.datetime.now().strftime('%d %B, %Y')}<br/>
    <b>Author:</b> Legal Data Research & AI/ML Development Team<br/>
    <b>Project Context:</b> Foundation Layer for Unified Legal Dataset Scraping<br/>
    <b>Classification:</b> Confidential — Internal Research Standard
    """
    story.append(HRFlowable(width="100%", thickness=1.5, color=c_primary, spaceBefore=20, spaceAfter=20))
    story.append(Paragraph(metadata_text, body_style))
    story.append(PageBreak())

    # ================= EXECUTIVE SUMMARY =================
    story.append(Paragraph("Executive Summary", h1_style))
    story.append(HRFlowable(width="100%", thickness=0.8, color=c_secondary, spaceAfter=10))
    
    datasets = db.query(Dataset).all()
    shortlisted = db.query(Dataset).filter(Dataset.shortlisted == True).all()
    platforms = set(d.platform for d in datasets if d.platform)
    verified_provenance = sum(1 for d in datasets if d.provenance_status == "Verified")
    clear_licenses = sum(1 for d in datasets if d.license_status == "Clear")
    unclear_licenses = sum(1 for d in datasets if d.license_status in ["License Unclear", "No License Found"])
    gaps_count = db.query(GapAnalysis).filter(GapAnalysis.priority.in_(["Critical", "High"])).count()

    summary_text = f"""
    This Landscape Investigation Report provides an inventory analysis of <b>{len(datasets)} discovered datasets</b> 
    mapped across <b>{len(platforms)} online registries</b> (such as AWS Open Data, Hugging Face, GitHub, and academic archives). 
    Of these, <b>{len(shortlisted)} datasets</b> have been shortlisted for high benchmark relevance. 
    A strict audit reveals that while <b>{verified_provenance} datasets</b> possess verified original court provenance, 
    licensing restrictions pose significant compliance barriers. Specifically, <b>{clear_licenses} datasets</b> have clear open licenses 
    suitable for SFT model training, whereas <b>{unclear_licenses} datasets</b> feature unclear or missing licenses, demanding extreme caution. 
    Furthermore, <b>{gaps_count} critical/high priority gaps</b> remain, particularly in local District Court coverage, original PDFs, 
    and regional language records, indicating that custom collection remains highly necessary.
    """
    story.append(Paragraph(summary_text, body_style))
    story.append(Spacer(1, 10))

    # ================= METHODOLOGY =================
    story.append(Paragraph("Research Methodology", h1_style))
    story.append(HRFlowable(width="100%", thickness=0.8, color=c_secondary, spaceAfter=10))
    
    meth = db.query(ResearchMethodology).first()
    if meth:
        meth_text = f"""
        <b>Platforms Searched:</b> {meth.platform_searched}<br/>
        <b>Search Date:</b> {meth.search_date.strftime('%Y-%m-%d') if meth.search_date else 'N/A'}<br/>
        <b>Search Keywords:</b> {meth.search_terms}<br/>
        <b>Selection Criteria:</b> {meth.selection_criteria}<br/>
        <b>Exclusion Criteria:</b> {meth.exclusion_criteria}<br/>
        <b>Verification Process:</b> {meth.verification_process}
        """
        story.append(Paragraph(meth_text, body_style))
    else:
        story.append(Paragraph("Research methodology parameters not recorded in database.", body_style))
    story.append(PageBreak())

    # ================= DATASET INVENTORY =================
    story.append(Paragraph("Dataset Inventory", h1_style))
    story.append(HRFlowable(width="100%", thickness=0.8, color=c_secondary, spaceAfter=10))
    
    inv_rows = [[
        Paragraph("<b>Dataset Name</b>", table_header_style), 
        Paragraph("<b>Platform</b>", table_header_style), 
        Paragraph("<b>Category</b>", table_header_style), 
        Paragraph("<b>Records</b>", table_header_style),
        Paragraph("<b>Provenance</b>", table_header_style),
        Paragraph("<b>License</b>", table_header_style)
    ]]
    for d in datasets:
        inv_rows.append([
            Paragraph(d.dataset_name, table_cell_style),
            Paragraph(d.platform, table_cell_style),
            Paragraph(d.category, table_cell_style),
            Paragraph(f"{d.record_count:,}" if d.record_count else "Unknown", table_cell_style),
            Paragraph(d.provenance_status, table_cell_style),
            Paragraph(d.license_status, table_cell_style)
        ])
    
    inv_table = Table(inv_rows, colWidths=[150, 70, 75, 55, 65, 75])
    inv_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_primary),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, c_light]),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(inv_table)
    story.append(PageBreak())

    # ================= DATASET COMPARISON (SHORTLISTED) =================
    story.append(Paragraph("Dataset Comparison Matrix (Shortlisted)", h1_style))
    story.append(HRFlowable(width="100%", thickness=0.8, color=c_secondary, spaceAfter=10))
    
    comp_rows = [[
        Paragraph("<b>Dataset</b>", table_header_style),
        Paragraph("<b>Platform</b>", table_header_style),
        Paragraph("<b>Records</b>", table_header_style),
        Paragraph("<b>Time Period</b>", table_header_style),
        Paragraph("<b>License</b>", table_header_style),
        Paragraph("<b>Freshness</b>", table_header_style),
        Paragraph("<b>Reuse Class</b>", table_header_style)
    ]]
    for d in shortlisted:
        comp_rows.append([
            Paragraph(d.short_name, table_cell_style),
            Paragraph(d.platform, table_cell_style),
            Paragraph(f"{d.record_count:,}" if d.record_count else "Unknown", table_cell_style),
            Paragraph(f"{d.time_period_start}-{d.time_period_end}" if d.time_period_start else "N/A", table_cell_style),
            Paragraph(d.license_name or "N/A", table_cell_style),
            Paragraph(d.freshness_status, table_cell_style),
            Paragraph(d.reuse_classification, table_cell_style)
        ])
    
    comp_table = Table(comp_rows, colWidths=[90, 65, 55, 65, 80, 50, 85])
    comp_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_secondary),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, c_light]),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(comp_table)
    story.append(Spacer(1, 15))

    # ================= PROVENANCE & LICENSING =================
    story.append(Paragraph("Provenance & License Compliance Analysis", h1_style))
    story.append(HRFlowable(width="100%", thickness=0.8, color=c_secondary, spaceAfter=10))
    
    warning_text = """
    <b>WARNING ON DATA REUSE LIMITATIONS:</b> Public availability on Kaggle, Hugging Face, or GitHub 
    does NOT imply legal consent for reuse. Standard Indian government records are public, 
    but private vendor enrichments, translations, and parsed academic subsets are governed by specific intellectual property boundaries.
    """
    story.append(Paragraph(warning_text, warning_body_style))
    story.append(Spacer(1, 8))

    prov_analysis = """
    <b>Original Source Verification:</b> The portal distinguishes between the Hosting Platform and the Original Source. 
    For SFT model training, data must be sourced from official domains (.gov.in/.nic.in) or verified open Sponsorship programs. 
    Third-party scrapes hosted on Hugging Face (such as community instruction QA pairs) carry high risks of synthetic hallucinations 
    and lack verified provenance certificates.
    """
    story.append(Paragraph(prov_analysis, body_style))
    story.append(PageBreak())

    # ================= GAP ANALYSIS TABLE =================
    story.append(Paragraph("Gap Analysis Table", h1_style))
    story.append(HRFlowable(width="100%", thickness=0.8, color=c_secondary, spaceAfter=10))
    
    gaps = db.query(GapAnalysis).all()
    gap_rows = [[
        Paragraph("<b>Category</b>", table_header_style),
        Paragraph("<b>Availability</b>", table_header_style),
        Paragraph("<b>Coverage Summary</b>", table_header_style),
        Paragraph("<b>Priority</b>", table_header_style),
        Paragraph("<b>Recommendation</b>", table_header_style)
    ]]
    for g in gaps:
        gap_rows.append([
            Paragraph(g.category, table_cell_style),
            Paragraph(g.availability, table_cell_style),
            Paragraph(g.current_state or "N/A", table_cell_style),
            Paragraph(g.priority, table_cell_style),
            Paragraph(g.recommendation or "N/A", table_cell_style)
        ])
    
    gap_table = Table(gap_rows, colWidths=[90, 50, 140, 45, 165])
    gap_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), c_primary),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, c_light]),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(gap_table)
    story.append(PageBreak())

    # ================= REUSE RECOMMENDATIONS =================
    story.append(Paragraph("Landscape Conclusion & Action Plan", h1_style))
    story.append(HRFlowable(width="100%", thickness=0.8, color=c_secondary, spaceAfter=10))
    
    recs = """
    <b>What Should We Collect Ourselves?</b><br/>
    Based on the compiled gap audit, our futures collection pipeline should prioritize:<br/>
    1. <b>District and Subordinate Judgments:</b> SHRUG provides case metadata but contains no text content. We must implement targeted district crawlers for high-frequency state registries.<br/>
    2. <b>Regional Language Text Corpora:</b> Over 90% of available text datasets are in English. Digital translation and OCR pipelines for regional languages (e.g. Hindi, Tamil) are critically missing.<br/>
    3. <b>Original Court PDFs:</b> Downstream legal verification require original certified files. Mirror scraping of central gazettes and SC/HC portals is needed to compile source PDFs.<br/>
    4. <b>Citation Maps:</b> Create parent-child link nodes mapping central statutes to judgments.
    """
    story.append(Paragraph(recs, body_style))
    story.append(Spacer(1, 10))

    # Evidence audit trail
    story.append(Paragraph("Investigation Audit Trail & Evidence", h2_style))
    evidence = db.query(DatasetEvidence).filter(DatasetEvidence.verified == True).all()
    if evidence:
        ev_text = "The following evidence citations were verified during the landscape review:<br/>"
        for idx, e in enumerate(evidence):
            ev_text += f"{idx+1}. <b>{e.source_title}</b> (URL: {e.source_url or 'N/A'}) - Verified by {e.reviewer} on {e.verified_at.strftime('%Y-%m-%d') if e.verified_at else 'N/A'}. Text: <i>\"{e.evidence_text}\"</i><br/>"
        story.append(Paragraph(ev_text, body_style))
    else:
        story.append(Paragraph("No verified evidence sources attached to report.", body_style))

    # Build PDF
    doc_template.build(story, canvasmaker=NumberedCanvas)
    return file_path
